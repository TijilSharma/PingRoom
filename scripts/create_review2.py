from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path('output/PingRoom_Review2_Database_Implementation.docx')
OUT.parent.mkdir(exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10)
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
for name, size, color in [('Title', 25, '1F3C31'), ('Heading 1', 16, '1F3C31'), ('Heading 2', 12, '385A46')]:
    st = styles[name]
    st.font.name = 'Aptos Display'
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
styles['Heading 1'].paragraph_format.space_before = Pt(16)
styles['Heading 1'].paragraph_format.space_after = Pt(7)
styles['Heading 2'].paragraph_format.space_before = Pt(11)
styles['Heading 2'].paragraph_format.space_after = Pt(5)
title_pPr = styles['Title']._element.get_or_add_pPr()
for border in title_pPr.findall(qn('w:pBdr')):
    title_pPr.remove(border)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    node = OxmlElement('w:shd')
    node.set(qn('w:fill'), fill)
    tcPr.append(node)

def set_cell_text(cell, value, bold=False, color=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(value)
    run.bold = bold
    run.font.size = Pt(8.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], h, True, 'FFFFFF')
        shade(t.rows[0].cells[i], '294638')
        if widths: t.rows[0].cells[i].width = Inches(widths[i])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell_text(cells[i], str(v), i == 0)
            if len(t.rows) % 2 == 1: shade(cells[i], 'F1F5EA')
            if widths: cells[i].width = Inches(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def para(text='', bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)

# Title page
p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run('PingRoom Review 2 Database Implementation')
# LibreOffice inherits a visible title rule from some built-in Title styles.
# Remove the paragraph border so the title remains plain and formal.
pPr = p._p.get_or_add_pPr()
for border in pPr.findall(qn('w:pBdr')):
    pPr.remove(border)
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('MERN Real Time Chat Application')
r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string('59705B')
doc.add_paragraph()
meta = doc.add_table(rows=4, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate([
    ('Course review', 'Review 2 Database Implementation and Core Functionalities'),
    ('Application', 'PingRoom'),
    ('Database', 'MongoDB 8 replica set with Mongoose'),
    ('Prepared for', 'Database Systems evaluation'),
]):
    set_cell_text(meta.rows[i].cells[0], k, True, 'FFFFFF'); shade(meta.rows[i].cells[0], '294638')
    set_cell_text(meta.rows[i].cells[1], v); shade(meta.rows[i].cells[1], 'F1F5EA')
doc.add_paragraph()
intro = doc.add_paragraph()
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
rr = intro.add_run('Review conclusion: PingRoom meets the Review 2 requirements with a working MongoDB database, complete CRUD operations, 138 seeded records, and demonstrated indexing, aggregation, transactions, and real time messaging.')
rr.italic = True; rr.font.color.rgb = RGBColor.from_string('4B6551')

doc.add_page_break()
doc.add_heading('1 Review 2 Scope and Result', level=1)
para('PingRoom is a full stack chat application for direct and group conversations. This review covers the database layer and the core application functions requested for Review 2. The implementation uses MongoDB as a single replica set, Express and Node.js for the API, React for the interface, and Socket.IO for live events.')
table(['Review requirement', 'Implemented evidence', 'Result'], [
    ('Database implementation', 'Three MongoDB collections, validation schemas, references, compound and unique indexes.', 'Complete'),
    ('CRUD operations', 'Authenticated create, read, update, and delete endpoints for users, conversations, and messages.', 'Complete'),
    ('Advanced NoSQL features', 'Aggregation pipelines, multikey/text indexes, execution-plan inspection, replica-set transactions.', 'Complete'),
    ('Sample data', '12 users, 6 conversations, and 120 messages.', '138 total records'),
    ('Intermediate demonstration', 'Database Studio displays counts, indexes, aggregation output, replica-set status, and query statistics.', 'Ready'),
    ('Source repository', 'Repository contains source code, Docker setup, seed script, automated tests, and this report.', 'Ready'),
], [1.35, 4.3, 1.0])

doc.add_page_break()
doc.add_heading('2 Database Design', level=1)
para('The design keeps shared users in a users collection, stores the participant list directly in each conversation, and stores messages separately. Messages are not embedded in conversations because a conversation can grow without a fixed limit. Each conversation maintains a lastMessage snapshot for fast conversation-list rendering.')
table(['Collection', 'Key fields', 'Relationship and purpose'], [
    ('users', '_id, username, email, passwordHash, status, lastSeen', 'One user can participate in many conversations and author many messages.'),
    ('conversations', '_id, type, title, ownerId, participants, lastMessage, updatedAt', 'Participants are embedded ObjectId references. Supports direct and group conversations.'),
    ('messages', '_id, conversationId, senderId, text, attachment, status, readBy, deliveredTo, createdAt', 'Each message references one conversation and one sender. Receipt arrays support group delivery and read status.'),
], [1.25, 2.8, 2.6])
doc.add_heading('Indexes', level=2)
bullet('users.username and users.email use unique indexes for account identity and login.')
bullet('conversations.participants with updatedAt is a multikey compound index for the conversation list ordered by recent activity.')
bullet('messages.conversationId with createdAt and _id is a compound index for newest-first history and stable cursor pagination.')
bullet('messages.text uses a text index for in-conversation message search.')

doc.add_page_break()
doc.add_heading('3 CRUD Operations', level=1)
table(['Entity', 'Create', 'Read', 'Update', 'Delete'], [
    ('User', 'Register with bcrypt hash', 'Current user and user search', 'Change username', 'Password-confirmed account deletion'),
    ('Conversation', 'Direct chat or group', 'Recent conversations and unread count', 'Group title by owner', 'Owner deletes conversation and all messages'),
    ('Message', 'Text or validated attachment link', 'Searchable paginated history', 'Author edits text or attachment', 'Author deletes message'),
], [1.15, 1.75, 1.8, 1.75, 1.75])
para('All protected API routes require a signed JWT. Request data is validated before writing to the database. The server also verifies membership before it permits access to a conversation, and verifies the owner or sender before it permits rename, edit, or delete operations.')

doc.add_heading('4 Advanced MongoDB Features', level=1)
doc.add_heading('Aggregation', level=2)
para('The conversation list computes unread messages with an aggregation pipeline. It matches messages belonging to the current user’s conversations, excludes messages sent by that user, filters out messages already read by that user, then groups the remaining records by conversation. The Database Studio also groups message counts and character totals by conversation, joins conversation titles through $lookup, and groups message volume by day.')
doc.add_heading('Transactions', level=2)
para('MongoDB transactions protect multi-document operations. Sending, editing, or deleting a message updates the message document and the conversation lastMessage snapshot in one transaction. Deleting a conversation removes its messages in the same transaction. Account deletion also cleans references and messages transactionally. The Docker configuration starts MongoDB as replica set rs0 because multi-document transactions require a replica set or sharded deployment.')
doc.add_heading('Query efficiency', level=2)
para('The Database Studio runs explain with execution statistics for a recent-message query. It exposes the selected query plan, documents examined, keys examined, and returned rows, allowing the reviewer to verify that the compound message-history index is in use.')

doc.add_page_break()
doc.add_heading('5 Sample Dataset and Demonstration Flow', level=1)
table(['Data set', 'Count', 'Notes'], [
    ('Users', '12', 'Includes the demo account tijil and eleven additional sample accounts.'),
    ('Conversations', '6', 'Three groups and three direct conversations.'),
    ('Messages', '120', 'Twenty messages per conversation with varied receipts and timestamps.'),
    ('Total records', '138', 'Exceeds the minimum 50 sample-record requirement.'),
], [1.6, .9, 5.7])
doc.add_heading('Suggested Review Demonstration', level=2)
for item in [
    'Open the application at http://localhost:3100 and sign in with tijil@vitstudent.ac.in using password PingRoom@123.',
    'Open a seeded conversation and scroll upward to demonstrate cursor-based message history.',
    'Send a message, edit it, and delete it. The conversation preview updates with each operation.',
    'Create a group conversation, then use a second browser session with a sample account to demonstrate live delivery, typing, presence, and receipts.',
    'Open Database Studio and show the collection counts, aggregation panels, index list, replica set rs0, and query execution statistics.',
    'Run npm test to show the automated verification suite.',
]: bullet(item)

doc.add_heading('6 Validation', level=1)
para('The automated test suite completed successfully with 10 passing tests. It verifies seeded data, password hashing, duplicate-account rejection, authentication, authorization, direct-chat de-duplication, message CRUD, receipt progression, stable pagination, group permissions, transaction rollback, aggregation scoping, index selection, Socket.IO authentication, typing, live delivery, multi-tab presence, and account deletion cleanup.')
table(['Command', 'Purpose', 'Observed result'], [
    ('npm run seed', 'Creates the sample dataset without deleting existing data.', '138 records created on an empty database.'),
    ('npm run build', 'Builds the React production client.', 'Build completed successfully.'),
    ('npm test', 'Runs API, transaction, database, and WebSocket tests.', '10 tests passed, 0 failed.'),
], [1.45, 3.8, 2.95])

doc.add_heading('7 Repository Contents', level=1)
para('The repository contains the React client, Express API, Mongoose models, Socket.IO integration, Docker Compose MongoDB replica-set configuration, seed script, setup script, test suite, environment template, and this formal Review 2 report. The project can be recreated with npm install, npm run setup, and npm start.')
final = doc.add_paragraph()
final.paragraph_format.space_before = Pt(15)
final.paragraph_format.space_after = Pt(0)
run = final.add_run('Review 2 status: Ready for evaluation.')
run.bold = True; run.font.color.rgb = RGBColor.from_string('294638')

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('PingRoom  |  Review 2 Database Implementation')
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string('8A9882')

doc.core_properties.title = 'PingRoom Review 2 Database Implementation'
doc.core_properties.author = 'Tijil Sharma'
doc.save(OUT)
print(OUT)
